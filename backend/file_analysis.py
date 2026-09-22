import io

from fastapi import APIRouter, HTTPException, UploadFile, File, Form

from ai import ask_eon


MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_EXTRACTED_CHARS = 120_000


def _ext(name: str) -> str:
    return (name.rsplit(".", 1)[-1].lower() if "." in name else "")


async def _read_upload(upload: UploadFile) -> bytes:
    data = await upload.read()
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail=f"{upload.filename} exceeds the 15 MB per-file limit.")
    return data


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "".join(
        f"\n--- PAGE {i + 1} ---\n{page.extract_text() or ''}"
        for i, page in enumerate(reader.pages)
    )


def _extract_docx(data: bytes) -> str:
    from docx import Document
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for i, table in enumerate(document.tables):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        parts.append(f"\n--- TABLE {i + 1} ---\n" + "\n".join(rows))
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation
    presentation = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(presentation.slides):
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        slides.append(f"\n--- SLIDE {i + 1} ---\n" + "\n".join(texts))
    return "".join(slides)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                rows.append(" | ".join(values))
            if len(rows) >= 500:
                break
        sheets.append(f"\n--- SHEET {sheet.title} ---\n" + "\n".join(rows))
    return "".join(sheets)


def _extract_plain(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _extract(filename: str, data: bytes) -> str:
    ext = _ext(filename)
    if ext == "pdf":
        return _extract_pdf(data)
    if ext == "docx":
        return _extract_docx(data)
    if ext == "pptx":
        return _extract_pptx(data)
    if ext == "xlsx":
        return _extract_xlsx(data)
    if ext == "xls":
        raise ValueError("Legacy .xls files are not supported yet. Save it as .xlsx first.")
    if ext in {"csv", "txt", "md", "json"}:
        return _extract_plain(data)
    raise ValueError(f"Unsupported file type: .{ext or 'unknown'}")


async def analyze_uploaded_files(
    uploads: list[UploadFile],
    question: str,
    mode: str = "NORMAL",
):
    blocks = []

    for upload in uploads:
        data = await _read_upload(upload)
        try:
            text = _extract(upload.filename or "file", data)
        except ImportError as error:
            raise HTTPException(
                status_code=500,
                detail=f"File parser dependency is missing: {error.name}.",
            )
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error))
        except Exception as error:
            raise HTTPException(
                status_code=400,
                detail=f"Could not read {upload.filename}: {error}",
            )

        text = text.strip()
        if not text:
            text = "[No extractable text/data was found in this file.]"

        blocks.append(
            f"\n===== FILE: {upload.filename} =====\n{text[:MAX_EXTRACTED_CHARS]}"
        )

    combined = "".join(blocks)

    prompt = f"""
You are EON's File Intelligence engine.

Analyze the uploaded files below.

User request:
{question.strip() or "Provide an executive summary, key findings, important numbers, and useful relationships between the uploaded files."}

Rules:
- Ground the answer only in the extracted file contents.
- If information is missing, say so.
- Keep file names and section/page/sheet references when available.
- For spreadsheets, discuss visible data and calculations supported by the extracted rows.
- For multiple files, explicitly identify useful cross-file relationships.
- Do not claim to have seen images, formatting, charts, or scanned text that was not extracted.
- Clearly separate observations from uncertainty.

Uploaded file contents:
{combined}
"""

    result = await ask_eon(
        message=prompt,
        mode=mode,
        memory_context="",
        destination="AI",
    )

    return result
